from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.http import HttpResponse
from django.utils import timezone
from datetime import timedelta, date
from collections import OrderedDict

from atendimentos.models import Atendimento, Doenca
from castracoes.models import Castracao
from cadastros.models import Animal, Tutor, ObitoAnimal
from vacinacoes.models import BoletimVacinacaoRaiva
from visitas.models import Visita

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import io


@login_required
def dashboard(request):
    hoje = timezone.now().date()
    proximos_7_dias = hoje + timedelta(days=7)

    # --- Período selecionado (mês/ano de referência para os cartões e status) ---
    try:
        ano_ref = int(request.GET.get('ano', hoje.year))
        mes_ref = int(request.GET.get('mes', hoje.month))
        inicio_mes = date(ano_ref, mes_ref, 1)
    except (ValueError, TypeError):
        ano_ref, mes_ref = hoje.year, hoje.month
        inicio_mes = hoje.replace(day=1)

    if mes_ref == 12:
        fim_mes = date(ano_ref + 1, 1, 1) - timedelta(days=1)
    else:
        fim_mes = date(ano_ref, mes_ref + 1, 1) - timedelta(days=1)

    # --- Quantidade de meses no gráfico de histórico ---
    try:
        qtd_meses_historico = int(request.GET.get('meses', 6))
        if qtd_meses_historico not in (3, 6, 12):
            qtd_meses_historico = 6
    except (ValueError, TypeError):
        qtd_meses_historico = 6

    inicio_historico = (inicio_mes - timedelta(days=30 * qtd_meses_historico)).replace(day=1)

    atendimentos_hoje = Atendimento.objects.filter(data_atendimento__date=hoje).count()
    atendimentos_mes = Atendimento.objects.filter(
        data_atendimento__date__gte=inicio_mes, data_atendimento__date__lte=fim_mes
    ).count()
    castracoes_hoje = Castracao.objects.filter(data_agendada=hoje).count()
    castracoes_realizadas_mes = Castracao.objects.filter(
        data_agendada__gte=inicio_mes, data_agendada__lte=fim_mes, status='REALIZADA'
    ).count()
    animais_ativos = Animal.objects.filter(status='ATIVO').count()
    tutores_total = Tutor.objects.count()
    obitos_mes = ObitoAnimal.objects.filter(data_obito__gte=inicio_mes, data_obito__lte=fim_mes).count()
    visitas_mes = Visita.objects.filter(data_visita__gte=inicio_mes, data_visita__lte=fim_mes).count()

    boletins_mes = BoletimVacinacaoRaiva.objects.filter(data__gte=inicio_mes, data__lte=fim_mes)
    vacinacoes_mes = sum(b.total for b in boletins_mes)

    atendimentos_por_mes = OrderedDict()
    mes_atual = inicio_historico
    while mes_atual <= inicio_mes:
        chave = mes_atual.strftime('%m/%Y')
        atendimentos_por_mes[chave] = 0
        if mes_atual.month == 12:
            mes_atual = mes_atual.replace(year=mes_atual.year + 1, month=1)
        else:
            mes_atual = mes_atual.replace(month=mes_atual.month + 1)

    atendimentos_qs = Atendimento.objects.filter(data_atendimento__gte=inicio_historico)
    for at in atendimentos_qs:
        chave = at.data_atendimento.strftime('%m/%Y')
        if chave in atendimentos_por_mes:
            atendimentos_por_mes[chave] += 1

    castracoes_status = {
        'Agendada': Castracao.objects.filter(data_agendada__gte=inicio_mes, data_agendada__lte=fim_mes, status='AGENDADA').count(),
        'Compareceu': Castracao.objects.filter(data_agendada__gte=inicio_mes, data_agendada__lte=fim_mes, status='COMPARECEU').count(),
        'Faltou': Castracao.objects.filter(data_agendada__gte=inicio_mes, data_agendada__lte=fim_mes, status='FALTOU').count(),
        'Realizada': Castracao.objects.filter(data_agendada__gte=inicio_mes, data_agendada__lte=fim_mes, status='REALIZADA').count(),
        'Cancelada': Castracao.objects.filter(data_agendada__gte=inicio_mes, data_agendada__lte=fim_mes, status='CANCELADA').count(),
    }

    animais_especies = {
        'Cão': Animal.objects.filter(status='ATIVO', especie='CAO').count(),
        'Gato': Animal.objects.filter(status='ATIVO', especie='GATO').count(),
        'Outro': Animal.objects.filter(status='ATIVO', especie='OUTRO').count(),
    }

    doencas_ranking = (
        Atendimento.objects.filter(doenca_confirmada__isnull=False)
        .values('doenca_confirmada__nome')
        .annotate(total=Count('id'))
        .order_by('-total')[:5]
    )

    vacinacoes_tipos = {
        'Rotina': sum(b.total for b in BoletimVacinacaoRaiva.objects.filter(tipo_vacinacao='ROTINA')),
        'Campanha': sum(b.total for b in BoletimVacinacaoRaiva.objects.filter(tipo_vacinacao='CAMPANHA')),
        'Cobertura de foco': sum(b.total for b in BoletimVacinacaoRaiva.objects.filter(tipo_vacinacao='COBERTURA_FOCO')),
    }

    proximas_castracoes = Castracao.objects.filter(
        data_agendada__gte=hoje,
        data_agendada__lte=proximos_7_dias,
        status='AGENDADA'
    ).select_related('animal', 'tutor').order_by('data_agendada', 'periodo')[:10]

    ultimos_atendimentos = Atendimento.objects.select_related(
        'animal', 'veterinaria', 'doenca_confirmada'
    ).order_by('-data_atendimento')[:5]

    # --- Todas as faltas do mês selecionado (sem limite de 10) ---
    faltas_do_mes = Castracao.objects.filter(
        data_agendada__gte=inicio_mes, data_agendada__lte=fim_mes, status='FALTOU'
    ).select_related('animal', 'tutor').order_by('data_agendada', 'periodo')

    meses_nomes = [
        (1, 'Janeiro'), (2, 'Fevereiro'), (3, 'Março'), (4, 'Abril'),
        (5, 'Maio'), (6, 'Junho'), (7, 'Julho'), (8, 'Agosto'),
        (9, 'Setembro'), (10, 'Outubro'), (11, 'Novembro'), (12, 'Dezembro'),
    ]
    mes_nome_selecionado = dict(meses_nomes).get(mes_ref, '')

    contexto = {
        'atendimentos_mes': atendimentos_mes,
        'castracoes_hoje': castracoes_hoje,
        'castracoes_realizadas_mes': castracoes_realizadas_mes,
        'animais_ativos': animais_ativos,
        'tutores_total': tutores_total,
        'obitos_mes': obitos_mes,
        'vacinacoes_mes': vacinacoes_mes,
        'visitas_mes': visitas_mes,
        'atendimentos_por_mes_labels': list(atendimentos_por_mes.keys()),
        'atendimentos_por_mes_valores': list(atendimentos_por_mes.values()),
        'castracoes_status_labels': list(castracoes_status.keys()),
        'castracoes_status_valores': list(castracoes_status.values()),
        'animais_especies_labels': list(animais_especies.keys()),
        'animais_especies_valores': list(animais_especies.values()),
        'doencas_ranking_labels': [d['doenca_confirmada__nome'] for d in doencas_ranking],
        'doencas_ranking_valores': [d['total'] for d in doencas_ranking],
        'vacinacoes_tipos_labels': list(vacinacoes_tipos.keys()),
        'vacinacoes_tipos_valores': list(vacinacoes_tipos.values()),
        'proximas_castracoes': proximas_castracoes,
        'ultimos_atendimentos': ultimos_atendimentos,
        'faltas_do_mes': faltas_do_mes,
        'total_faltas_do_mes': faltas_do_mes.count(),
        'ano_atual': hoje.year,
        'anos_disponiveis': list(range(hoje.year, hoje.year - 5, -1)),
        # período selecionado, para o formulário de filtro
        'mes_selecionado': mes_ref,
        'ano_selecionado': ano_ref,
        'meses_historico_selecionado': qtd_meses_historico,
        'meses_nomes': meses_nomes,
        'mes_nome_selecionado': mes_nome_selecionado,
    }

    return render(request, 'relatorios/dashboard.html', contexto)


@login_required
def relatorio_geral(request):
    return render(request, 'relatorios/relatorio_geral.html', {})


MESES_NOMES = [
    '', 'Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho',
    'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro'
]


@login_required
def detalhes_castracoes(request):
    hoje = timezone.now().date()

    # Mês/ano escolhido pelo usuário via ?mes=8&ano=2026, com o mês atual como padrão
    try:
        mes_selecionado = int(request.GET.get('mes', hoje.month))
        ano_selecionado = int(request.GET.get('ano', hoje.year))
        if mes_selecionado < 1 or mes_selecionado > 12:
            raise ValueError
    except (TypeError, ValueError):
        mes_selecionado = hoje.month
        ano_selecionado = hoje.year

    inicio_mes = date(ano_selecionado, mes_selecionado, 1)
    if mes_selecionado == 12:
        fim_mes = date(ano_selecionado + 1, 1, 1)
    else:
        fim_mes = date(ano_selecionado, mes_selecionado + 1, 1)

    seis_meses_atras = (inicio_mes - timedelta(days=180)).replace(day=1)

    anos_disponiveis = list(range(hoje.year, hoje.year - 5, -1))

    castracoes_mes = Castracao.objects.filter(data_agendada__gte=inicio_mes, data_agendada__lt=fim_mes)
    total_mes = castracoes_mes.count()
    agendadas_mes = castracoes_mes.filter(status='AGENDADA').count()
    compareceram_mes = castracoes_mes.filter(status='COMPARECEU').count()
    faltaram_mes = castracoes_mes.filter(status='FALTOU').count()
    realizadas_mes = castracoes_mes.filter(status='REALIZADA').count()
    canceladas_mes = castracoes_mes.filter(status='CANCELADA').count()

    com_resultado = compareceram_mes + faltaram_mes + realizadas_mes
    if com_resultado > 0:
        taxa_comparecimento = round(
            ((compareceram_mes + realizadas_mes) / com_resultado) * 100, 1
        )
    else:
        taxa_comparecimento = 0

    castracoes_realizadas = castracoes_mes.filter(status__in=['REALIZADA', 'COMPARECEU']).select_related('animal')
    cao_macho = castracoes_realizadas.filter(animal__especie='CAO', animal__sexo='MACHO').count()
    cao_femea = castracoes_realizadas.filter(animal__especie='CAO', animal__sexo='FEMEA').count()
    gato_macho = castracoes_realizadas.filter(animal__especie='GATO', animal__sexo='MACHO').count()
    gato_femea = castracoes_realizadas.filter(animal__especie='GATO', animal__sexo='FEMEA').count()

    manha_mes = castracoes_realizadas.filter(periodo='MANHA').count()
    tarde_mes = castracoes_realizadas.filter(periodo='TARDE').count()

    historico = OrderedDict()
    mes_atual = seis_meses_atras
    while mes_atual <= hoje:
        chave = mes_atual.strftime('%m/%Y')
        historico[chave] = {'realizadas': 0, 'faltas': 0}
        if mes_atual.month == 12:
            mes_atual = mes_atual.replace(year=mes_atual.year + 1, month=1)
        else:
            mes_atual = mes_atual.replace(month=mes_atual.month + 1)

    castracoes_periodo = Castracao.objects.filter(data_agendada__gte=seis_meses_atras)
    for c in castracoes_periodo:
        chave = c.data_agendada.strftime('%m/%Y')
        if chave in historico:
            if c.status in ['REALIZADA', 'COMPARECEU']:
                historico[chave]['realizadas'] += 1
            elif c.status == 'FALTOU':
                historico[chave]['faltas'] += 1

    faltas_do_mes = Castracao.objects.filter(
        data_agendada__gte=inicio_mes, data_agendada__lt=fim_mes, status='FALTOU'
    ).select_related('animal', 'tutor').order_by('data_agendada', 'periodo')

    proximas = Castracao.objects.filter(
        data_agendada__gte=hoje,
        status='AGENDADA'
    ).select_related('animal', 'tutor').order_by('data_agendada', 'periodo')[:15]

    meses_nomes_pt = {
        1: 'Janeiro', 2: 'Fevereiro', 3: 'Março', 4: 'Abril', 5: 'Maio', 6: 'Junho',
        7: 'Julho', 8: 'Agosto', 9: 'Setembro', 10: 'Outubro', 11: 'Novembro', 12: 'Dezembro',
    }

    contexto = {
        'total_mes': total_mes,
        'agendadas_mes': agendadas_mes,
        'compareceram_mes': compareceram_mes,
        'faltaram_mes': faltaram_mes,
        'realizadas_mes': realizadas_mes,
        'canceladas_mes': canceladas_mes,
        'taxa_comparecimento': taxa_comparecimento,
        'cao_macho': cao_macho,
        'cao_femea': cao_femea,
        'gato_macho': gato_macho,
        'gato_femea': gato_femea,
        'manha_mes': manha_mes,
        'tarde_mes': tarde_mes,
        'historico_labels': list(historico.keys()),
        'historico_realizadas': [h['realizadas'] for h in historico.values()],
        'historico_faltas': [h['faltas'] for h in historico.values()],
        'faltas_do_mes': faltas_do_mes,
        'total_faltas_do_mes': faltas_do_mes.count(),
        'proximas': proximas,
        'mes_selecionado': mes_selecionado,
        'ano_selecionado': ano_selecionado,
        'mes_nome_selecionado': meses_nomes_pt.get(mes_selecionado, ''),
        'anos_disponiveis': anos_disponiveis,
        'meses_nomes': sorted(meses_nomes_pt.items()),
    }

    return render(request, 'relatorios/detalhes_castracoes.html', contexto)


MESES_PT = [
    'JANEIRO', 'FEVEREIRO', 'MARÇO', 'ABRIL',
    'MAIO', 'JUNHO', 'JULHO', 'AGOSTO',
    'SETEMBRO', 'OUTUBRO', 'NOVEMBRO', 'DEZEMBRO',
]


def _coleta_dados_mes(ano, mes):
    primeiro_dia = date(ano, mes, 1)
    if mes == 12:
        ultimo_dia = date(ano + 1, 1, 1) - timedelta(days=1)
    else:
        ultimo_dia = date(ano, mes + 1, 1) - timedelta(days=1)

    visitas_urbana = Visita.objects.filter(
        data_visita__gte=primeiro_dia, data_visita__lte=ultimo_dia, area='URBANA'
    ).count()
    visitas_rural = Visita.objects.filter(
        data_visita__gte=primeiro_dia, data_visita__lte=ultimo_dia, area='RURAL'
    ).count()

    castracoes_realizadas = Castracao.objects.filter(
        data_agendada__gte=primeiro_dia,
        data_agendada__lte=ultimo_dia,
        status='REALIZADA',
    ).select_related('animal')

    cao_macho = castracoes_realizadas.filter(animal__especie='CAO', animal__sexo='MACHO').count()
    cao_femea = castracoes_realizadas.filter(animal__especie='CAO', animal__sexo='FEMEA').count()
    gato_macho = castracoes_realizadas.filter(animal__especie='GATO', animal__sexo='MACHO').count()
    gato_femea = castracoes_realizadas.filter(animal__especie='GATO', animal__sexo='FEMEA').count()

    boletins_rotina = BoletimVacinacaoRaiva.objects.filter(
        data__gte=primeiro_dia, data__lte=ultimo_dia, tipo_vacinacao='ROTINA'
    )
    vacinacao_rotina = sum(b.total for b in boletins_rotina)

    animais_atendidos = Atendimento.objects.filter(
        data_atendimento__date__gte=primeiro_dia,
        data_atendimento__date__lte=ultimo_dia,
    ).count()

    total_visitas = visitas_urbana + visitas_rural
    total_animais = cao_macho + cao_femea + gato_macho + gato_femea + vacinacao_rotina + animais_atendidos

    return {
        'visitas_urbana': visitas_urbana,
        'visitas_rural': visitas_rural,
        'cao_macho': cao_macho,
        'cao_femea': cao_femea,
        'gato_macho': gato_macho,
        'gato_femea': gato_femea,
        'vacinacao_rotina': vacinacao_rotina,
        'teste_lva': '',
        'animais_atendidos': animais_atendidos,
        'suporte_pm': '',
        'total_visitas': total_visitas,
        'total_animais': total_animais,
    }


def _formatar(v):
    if v == '' or v is None:
        return '-'
    if isinstance(v, int) and v == 0:
        return '0'
    if isinstance(v, int) and v < 10:
        return f'0{v}'
    return str(v)


def _set_cell_text(cell, text, bold=True, size=9, center=True):
    cell.text = ''
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _sombrear_celula(cell, cor_hex='D9D9D9'):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), cor_hex)
    tc_pr.append(shd)


@login_required
def gerar_resumo_anual(request):
    ano = int(request.GET.get('ano', timezone.now().year))

    dados_meses = {}
    for mes in range(1, 13):
        dados_meses[mes] = _coleta_dados_mes(ano, mes)

    def soma_quadrimestre(meses):
        chaves = ['visitas_urbana', 'visitas_rural', 'cao_macho', 'cao_femea',
                  'gato_macho', 'gato_femea', 'vacinacao_rotina', 'animais_atendidos',
                  'total_visitas', 'total_animais']
        resultado = {}
        for chave in chaves:
            resultado[chave] = sum(dados_meses[m][chave] for m in meses)
        resultado['teste_lva'] = ''
        resultado['suporte_pm'] = ''
        return resultado

    quad1 = soma_quadrimestre([1, 2, 3, 4])
    quad2 = soma_quadrimestre([5, 6, 7, 8])
    quad3 = soma_quadrimestre([9, 10, 11, 12])

    anual_total_animais = sum([
        quad1['cao_macho'], quad1['cao_femea'], quad1['gato_macho'], quad1['gato_femea'],
        quad2['cao_macho'], quad2['cao_femea'], quad2['gato_macho'], quad2['gato_femea'],
        quad3['cao_macho'], quad3['cao_femea'], quad3['gato_macho'], quad3['gato_femea'],
    ])

    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(1)
        section.bottom_margin = Cm(1)
        section.left_margin = Cm(1)
        section.right_margin = Cm(1)

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run('CENTRO DE CONTROLE DE ZOONOSES/CASA PET')
    run.bold = True
    run.font.size = Pt(12)

    subtitulo = doc.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitulo.add_run(f'RESUMO DE ATIVIDADES – {ano}')
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph()

    sequencia = []
    for i, mes_nome in enumerate(MESES_PT):
        sequencia.append(('mes', i + 1, mes_nome))
        sequencia.append(('total_mes', i + 1, None))
        if (i + 1) in [4, 8, 12]:
            quad_num = (i + 1) // 4
            sequencia.append(('quadrimestre', quad_num, f'{quad_num}º QUADRIM.'))
            sequencia.append(('total_quad', quad_num, None))
    sequencia.append(('anual', 0, 'ANUAL'))
    sequencia.append(('total_anual', 0, None))

    n_colunas = 11
    n_linhas = 2 + len(sequencia)

    tabela = doc.add_table(rows=n_linhas, cols=n_colunas)
    tabela.style = 'Table Grid'

    cab1 = tabela.rows[0]
    _set_cell_text(cab1.cells[0], 'ATIVIDADE/MÊS', size=8)
    _sombrear_celula(cab1.cells[0])
    _set_cell_text(cab1.cells[1], 'VISITAS ZONA URBANA', size=8)
    _sombrear_celula(cab1.cells[1])
    _set_cell_text(cab1.cells[2], 'VISITAS ZONA RURAL', size=8)
    _sombrear_celula(cab1.cells[2])

    cao_merged = cab1.cells[3].merge(cab1.cells[4])
    _set_cell_text(cao_merged, 'CASTRAÇÃO CANINA', size=8)
    _sombrear_celula(cao_merged)

    gato_merged = cab1.cells[5].merge(cab1.cells[6])
    _set_cell_text(gato_merged, 'CASTRAÇÃO FELINA', size=8)
    _sombrear_celula(gato_merged)

    _set_cell_text(cab1.cells[7], 'VACINAÇÃO ANTIRRÁBICA ROTINA', size=8)
    _sombrear_celula(cab1.cells[7])
    _set_cell_text(cab1.cells[8], 'TESTE RÁPIDO DE LVA', size=8)
    _sombrear_celula(cab1.cells[8])
    _set_cell_text(cab1.cells[9], 'ANIMAIS ATENDIDOS', size=8)
    _sombrear_celula(cab1.cells[9])
    _set_cell_text(cab1.cells[10], 'SUPORTE POLÍCIA AMBIENTAL', size=8)
    _sombrear_celula(cab1.cells[10])

    cab2 = tabela.rows[1]
    _set_cell_text(cab2.cells[0], '', size=8)
    _set_cell_text(cab2.cells[1], '', size=8)
    _set_cell_text(cab2.cells[2], '', size=8)
    _set_cell_text(cab2.cells[3], 'MACHO', size=8)
    _sombrear_celula(cab2.cells[3])
    _set_cell_text(cab2.cells[4], 'FÊMEA', size=8)
    _sombrear_celula(cab2.cells[4])
    _set_cell_text(cab2.cells[5], 'MACHO', size=8)
    _sombrear_celula(cab2.cells[5])
    _set_cell_text(cab2.cells[6], 'FÊMEA', size=8)
    _sombrear_celula(cab2.cells[6])
    for i in range(7, 11):
        _set_cell_text(cab2.cells[i], '', size=8)

    linha_idx = 2
    for tipo, num, nome in sequencia:
        row = tabela.rows[linha_idx]

        if tipo == 'mes':
            d = dados_meses[num]
            _set_cell_text(row.cells[0], nome, size=8)
            _sombrear_celula(row.cells[0])
            _set_cell_text(row.cells[1], _formatar(d['visitas_urbana']), bold=True, size=8)
            _set_cell_text(row.cells[2], _formatar(d['visitas_rural']), bold=True, size=8)
            _set_cell_text(row.cells[3], _formatar(d['cao_macho']), bold=True, size=8)
            _set_cell_text(row.cells[4], _formatar(d['cao_femea']), bold=True, size=8)
            _set_cell_text(row.cells[5], _formatar(d['gato_macho']), bold=True, size=8)
            _set_cell_text(row.cells[6], _formatar(d['gato_femea']), bold=True, size=8)
            _set_cell_text(row.cells[7], _formatar(d['vacinacao_rotina']), bold=True, size=8)
            _set_cell_text(row.cells[8], '-', bold=True, size=8)
            _set_cell_text(row.cells[9], _formatar(d['animais_atendidos']), bold=True, size=8)
            _set_cell_text(row.cells[10], '-', bold=True, size=8)

        elif tipo == 'total_mes':
            d = dados_meses[num]
            c0 = row.cells[0]
            _set_cell_text(c0, '', size=8)
            v_merged = row.cells[1].merge(row.cells[2])
            _set_cell_text(v_merged, f'TOTAL DE VISITAS: {_formatar(d["total_visitas"])}', size=8)
            a_merged = row.cells[3].merge(row.cells[10])
            _set_cell_text(a_merged, f'TOTAL DE ANIMAIS: {_formatar(d["total_animais"])}', size=8)

        elif tipo == 'quadrimestre':
            if num == 1:
                d = quad1
            elif num == 2:
                d = quad2
            else:
                d = quad3
            _set_cell_text(row.cells[0], nome, size=8)
            _sombrear_celula(row.cells[0])
            _set_cell_text(row.cells[1], _formatar(d['visitas_urbana']), bold=True, size=8)
            _set_cell_text(row.cells[2], _formatar(d['visitas_rural']), bold=True, size=8)
            _set_cell_text(row.cells[3], _formatar(d['cao_macho']), bold=True, size=8)
            _set_cell_text(row.cells[4], _formatar(d['cao_femea']), bold=True, size=8)
            _set_cell_text(row.cells[5], _formatar(d['gato_macho']), bold=True, size=8)
            _set_cell_text(row.cells[6], _formatar(d['gato_femea']), bold=True, size=8)
            _set_cell_text(row.cells[7], _formatar(d['vacinacao_rotina']), bold=True, size=8)
            _set_cell_text(row.cells[8], '-', bold=True, size=8)
            _set_cell_text(row.cells[9], _formatar(d['animais_atendidos']), bold=True, size=8)
            _set_cell_text(row.cells[10], '-', bold=True, size=8)
            for c in row.cells:
                _sombrear_celula(c, 'F2F2F2')

        elif tipo == 'total_quad':
            if num == 1:
                d = quad1
            elif num == 2:
                d = quad2
            else:
                d = quad3
            _set_cell_text(row.cells[0], '', size=8)
            v_merged = row.cells[1].merge(row.cells[2])
            _set_cell_text(v_merged, f'TOTAL DE VISITAS: {_formatar(d["total_visitas"])}', size=8)
            a_merged = row.cells[3].merge(row.cells[10])
            _set_cell_text(a_merged, f'TOTAL DE ANIMAIS: {_formatar(d["total_animais"])}', size=8)

        elif tipo == 'anual':
            _set_cell_text(row.cells[0], 'ANUAL', size=8)
            _sombrear_celula(row.cells[0], 'BFBFBF')
            cao_m = quad1['cao_macho'] + quad2['cao_macho'] + quad3['cao_macho']
            cao_f = quad1['cao_femea'] + quad2['cao_femea'] + quad3['cao_femea']
            gato_m = quad1['gato_macho'] + quad2['gato_macho'] + quad3['gato_macho']
            gato_f = quad1['gato_femea'] + quad2['gato_femea'] + quad3['gato_femea']
            _set_cell_text(row.cells[1], '', size=8)
            _set_cell_text(row.cells[2], '', size=8)
            _set_cell_text(row.cells[3], _formatar(cao_m), bold=True, size=8)
            _set_cell_text(row.cells[4], _formatar(cao_f), bold=True, size=8)
            _set_cell_text(row.cells[5], _formatar(gato_m), bold=True, size=8)
            _set_cell_text(row.cells[6], _formatar(gato_f), bold=True, size=8)
            _set_cell_text(row.cells[7], '', size=8)
            _set_cell_text(row.cells[8], '', size=8)
            _set_cell_text(row.cells[9], '', size=8)
            _set_cell_text(row.cells[10], '', size=8)
            for c in row.cells:
                _sombrear_celula(c, 'BFBFBF')

        elif tipo == 'total_anual':
            _set_cell_text(row.cells[0], '', size=8)
            v_merged = row.cells[1].merge(row.cells[2])
            total_v = quad1['total_visitas'] + quad2['total_visitas'] + quad3['total_visitas']
            _set_cell_text(v_merged, f'TOTAL DE VISITAS: {_formatar(total_v)}', size=8)
            a_merged = row.cells[3].merge(row.cells[10])
            _set_cell_text(a_merged, f'TOTAL DE ANIMAIS: {_formatar(anual_total_animais)}', size=8)

        linha_idx += 1

    doc.add_paragraph()
    titulo_cp = doc.add_paragraph()
    titulo_cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo_cp.add_run('CASA PET')
    run.bold = True
    run.font.size = Pt(12)

    tabela_cp = doc.add_table(rows=18, cols=7)
    tabela_cp.style = 'Table Grid'

    headers_cp = ['MÊS', 'Nº CÃES ABRIGADOS', 'Nº CÃES RECOLHIDOS', 'Nº CÃES DOADOS',
                  'Nº ÓBITOS', 'Nº EUTANÁSIAS', 'Nº CÃES ADOTADOS EM FEIRAS']
    cab = tabela_cp.rows[0]
    for i, h in enumerate(headers_cp):
        _set_cell_text(cab.cells[i], h, size=8)
        _sombrear_celula(cab.cells[i])

    linhas_cp = list(MESES_PT[:4]) + ['1º QUADRIM.'] + list(MESES_PT[4:8]) + ['2º QUADRIM.'] + \
                list(MESES_PT[8:12]) + ['3º QUADRIM.'] + ['ANUAL']
    for i, nome in enumerate(linhas_cp, start=1):
        row = tabela_cp.rows[i]
        _set_cell_text(row.cells[0], nome, size=8)
        if 'QUADRIM' in nome or nome == 'ANUAL':
            _sombrear_celula(row.cells[0], 'BFBFBF')
        else:
            _sombrear_celula(row.cells[0])
        for j in range(1, 7):
            _set_cell_text(row.cells[j], '', size=8)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = f'attachment; filename="resumo_anual_{ano}.docx"'
    return response